#!/usr/bin/env python3
"""Reproducible surface-form count and resumable per-word research store."""
import argparse
import csv
import hashlib
import json
import os
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
import re
import unicodedata

ROOT = Path(__file__).resolve().parent
CORPUS = ROOT / 'Quran_text.txt'
DATA = ROOT / 'data'
WORDS = DATA / 'words'
PROFILE = 'unvocalized_nfc_v1_numbered_verses'
VOCAL_PROFILE = 'vocalized_nfc_v1_numbered_verses'
BASMALA = 'بِسْمِ اللَّهِ الرَّحْمَـٰنِ الرَّحِيمِ'
DROP_MARKS = set(chr(i) for i in range(0x064B, 0x0653)) | {'\u0670'}
STOP_MARKS = set(chr(i) for i in range(0x06D6, 0x06DD))
DECORATION = {'\u06de', '\u06e9'}
REMOVED = DROP_MARKS | STOP_MARKS | DECORATION | {'\u0640'}
VERSE_RE = re.compile(r'^(\d+)\|(\d+)\|(.+)$')
STATUS_PENDING = 'Исследование не завершено'


def now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')


def digest(data):
    return hashlib.sha256(data).hexdigest()


def atomic_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    raw = (json.dumps(obj, ensure_ascii=False, sort_keys=True, indent=2) + '\n').encode()
    tmp = path.with_name(path.name + '.tmp')
    with tmp.open('wb') as handle:
        handle.write(raw)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(tmp, path)


def load_corpus():
    raw = CORPUS.read_bytes()
    body = raw.decode('utf-8', errors='strict')
    verses = []
    by_sura = defaultdict(list)
    chars = Counter()
    for line_no, line in enumerate(body.splitlines(), 1):
        match = VERSE_RE.fullmatch(line)
        if not match:
            raise ValueError(f'Invalid verse line {line_no}')
        sura, ayah = int(match[1]), int(match[2])
        text = match[3]
        if not text.strip():
            raise ValueError(f'Empty verse {sura}:{ayah}')
        chars.update(text)
        by_sura[sura].append(ayah)
        verses.append((sura, ayah, text))
    if len(verses) != 6236 or set(by_sura) != set(range(1, 115)):
        raise ValueError('Expected 6236 verses and suras 1–114')
    if len({(s, a) for s, a, _ in verses}) != len(verses):
        raise ValueError('Duplicate identifiers')
    if any(ayahs != list(range(1, len(ayahs) + 1)) for ayahs in by_sura.values()):
        raise ValueError('Nonconsecutive verse identifiers')
    if [(s, a) for s, a, _ in verses] != sorted((s, a) for s, a, _ in verses):
        raise ValueError('Verses out of order')
    unsupported = {c: n for c, n in chars.items() if not (c.isspace() or c in REMOVED or '\u0621' <= c <= '\u064a')}
    if unsupported:
        raise ValueError(f'Unclassified characters: {unsupported}')
    return raw, verses, by_sura, chars


def normalize(token):
    vocalized = unicodedata.normalize('NFC', ''.join(c for c in token if c not in STOP_MARKS | DECORATION | {'\u0640'}))
    plain = unicodedata.normalize('NFC', ''.join(c for c in vocalized if c not in DROP_MARKS))
    if not plain or not all('\u0621' <= c <= '\u064a' for c in plain):
        raise ValueError(f'Invalid normalized token {token!r} -> {plain!r}')
    return vocalized, plain


def corpus_rows(verses):
    basmala_plain = [normalize(t)[1] for t in BASMALA.split()]
    found = []
    records = []
    for sura, ayah, original in verses:
        text = original
        source_offset = 0
        if ayah == 1 and sura not in (1, 9):
            first = text.split(maxsplit=4)
            if len(first) != 5 or [normalize(t)[1] for t in first[:4]] != basmala_plain:
                raise ValueError(f'Missing initial basmala at {sura}:1')
            text = first[4]
            source_offset = 4
            found.append(sura)
        elif ayah == 1 and sura == 1 and text != BASMALA:
            raise ValueError('Al-Fatiha 1:1 differs from basmala')
        elif ayah == 1 and sura == 9 and [normalize(t)[1] for t in text.split()[:4]] == basmala_plain:
            raise ValueError('Unexpected 9:1 basmala')
        position = 0
        for source_pos, raw_token in enumerate(text.split(), 1):
            cleaned = ''.join(c for c in raw_token if c not in STOP_MARKS | DECORATION)
            if not cleaned:
                continue
            vocal, plain = normalize(raw_token)
            position += 1
            records.append((sura, ayah, position, source_pos + source_offset, raw_token, vocal, plain))
    if len(found) != 112:
        raise ValueError(f'Expected 112 unnumbered basmalas, found {len(found)}')
    return records, found


def identity(corpus_sha, key, profile=PROFILE):
    return digest((corpus_sha + '\0' + profile + '\0' + key).encode())


def build():
    raw, verses, by_sura, chars = load_corpus()
    sha = digest(raw)
    records, basmalas = corpus_rows(verses)
    counts = Counter(row[-1] for row in records)
    vocal_counts = Counter((row[-1], row[-2]) for row in records)
    independent = Counter()
    per_verse = defaultdict(int)
    per_sura = Counter()
    for sura, ayah, text in verses:
        if ayah == 1 and sura not in (1, 9):
            text = text.split(maxsplit=4)[4]
        # A separate streaming pass over verse text, without the occurrence index.
        for token in text.split():
            letters = unicodedata.normalize('NFC', ''.join(c for c in token if c not in REMOVED))
            if letters:
                independent[letters] += 1
    assert independent == counts
    for sura, ayah, *_ in records:
        per_verse[(sura, ayah)] += 1
        per_sura[sura] += 1
    assert sum(counts.values()) == len(records) == sum(per_verse.values()) == sum(per_sura.values())
    assert len(per_verse) == 6236
    previous = DATA / 'manifest.json'
    if previous.exists():
        old = json.loads(previous.read_text())
        if old['corpus_sha256'] != sha or old['profile_id'] != PROFILE:
            raise ValueError('Corpus or profile changed: create a new version; do not overwrite research')
    WORDS.mkdir(parents=True, exist_ok=True)
    occurrence_path = DATA / 'occurrences.csv'
    tmp = occurrence_path.with_name(occurrence_path.name + '.tmp')
    with tmp.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.writer(handle)
        writer.writerow(['word_id', 'sura', 'ayah', 'position', 'source_position', 'source_token', 'vocalized_nfc', 'word'])
        for sura, ayah, pos, source_pos, source, vocal, plain in records:
            writer.writerow([identity(sha, plain), sura, ayah, pos, source_pos, source, vocal, plain])
    os.replace(tmp, occurrence_path)
    freq_version = digest((''.join(f'{k}\t{v}\n' for k, v in sorted(counts.items())) + sha + PROFILE).encode())
    for key, frequency in counts.items():
        wid = identity(sha, key)
        path = WORDS / f'{wid}.json'
        if path.exists():
            current = json.loads(path.read_text())
            if (current['word_id'], current['word'], current['frequency'], current['frequency_version']) != (wid, key, frequency, freq_version):
                raise ValueError(f'Conflicting word record {path}')
            continue
        atomic_json(path, {
            'word_id': wid, 'profile_id': PROFILE, 'corpus_sha256': sha,
            'frequency_version': freq_version, 'word': key, 'frequency': frequency,
            'occurrences': f'data/occurrences.csv#word_id={wid}',
            'meanings': [], 'meaning_source': None, 'ambiguity': None,
            'research_status': STATUS_PENDING, 'hypotheses': [], 'external_facts': [],
            'sources': [], 'completed_steps': [], 'open_questions': [],
            'next_action': 'Проверить значение и контексты; затем применить протокол поиска.',
            'third_column': STATUS_PENDING, 'record_version': 1, 'updated_at': now()
        })
    with (DATA / 'vocalized_forms.csv').open('w', encoding='utf-8', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(['Словоформа с огласовками (NFC)', 'Количество', 'Словоформа без огласовок', 'word_id'])
        for (plain, vocal), n in sorted(vocal_counts.items(), key=lambda x: (-x[1], x[0])):
            writer.writerow([vocal, n, plain, identity(sha, plain)])
    with (DATA / 'unnumbered_initial_basmalas.csv').open('w', encoding='utf-8', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(['sura', 'source_ayah', 'source_position', 'source_token', 'vocalized_nfc', 'unvocalized_nfc'])
        for sura, ayah, source_text in verses:
            if ayah == 1 and sura in basmalas:
                for pos, token in enumerate(source_text.split()[:4], 1):
                    vocal, plain = normalize(token)
                    writer.writerow([sura, ayah, pos, token, vocal, plain])
    manifest = {
        'corpus_file': 'Quran_text.txt', 'corpus_sha256': sha, 'corpus_bytes': len(raw),
        'profile_id': PROFILE, 'vocalized_profile_id': VOCAL_PROFILE,
        'frequency_version': freq_version, 'numbered_verses': len(verses), 'suras': len(by_sura),
        'unnumbered_initial_basmalas_excluded': len(basmalas),
        'unnumbered_initial_basmala_tokens_excluded': 4 * len(basmalas),
        'unnumbered_initial_basmalas_suras': basmalas,
        'token_count': len(records), 'unique_words': len(counts),
        'unique_vocalized_forms': len(vocal_counts), 'generated_at': now(),
        'occurrences_sha256': digest(occurrence_path.read_bytes()),
        'independent_count_verified': True,
    }
    atomic_json(DATA / 'manifest.json', manifest)
    export()
    return manifest


def read_words():
    manifest = json.loads((DATA / 'manifest.json').read_text())
    if digest(CORPUS.read_bytes()) != manifest['corpus_sha256']:
        raise ValueError('Corpus hash mismatch')
    words = []
    for path in WORDS.glob('*.json'):
        item = json.loads(path.read_text())
        if item['word_id'] != path.stem or item['frequency_version'] != manifest['frequency_version']:
            raise ValueError(f'Invalid record {path}')
        words.append(item)
    if len(words) != manifest['unique_words'] or len({w['word_id'] for w in words}) != len(words):
        raise ValueError('Missing or duplicate word records')
    return manifest, words


def export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    manifest, words = read_words()
    words.sort(key=lambda w: (-w['frequency'], w['word']))
    rows = [['Слово', 'Количество', 'Связь с реальным миром']]
    rows.extend([w['word'], w['frequency'], w['third_column']] for w in words)
    output = DATA / 'quran_word_counts.csv'
    tmp = output.with_name(output.name + '.tmp')
    with tmp.open('w', encoding='utf-8-sig', newline='') as handle:
        csv.writer(handle).writerows(rows)
    os.replace(tmp, output)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Слова без огласовок'
    sheet.sheet_view.rightToLeft = True
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:C{sheet.max_row}'
    for col, width in [('A', 28), ('B', 15), ('C', 100)]:
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='254D64')
    for row in sheet.iter_rows(min_row=2):
        row[0].font = Font(name='Amiri', size=13)
        row[0].alignment = Alignment(horizontal='right', readingOrder=2)
        row[2].alignment = Alignment(wrap_text=True, vertical='top')
    out = DATA / 'quran_word_counts.xlsx'
    tmp = out.with_name(out.stem + '.tmp.xlsx')
    book.save(tmp)
    os.replace(tmp, out)
    status = Counter(w['research_status'] for w in words)
    queue = [{'word_id': w['word_id'], 'word': w['word'], 'frequency': w['frequency'], 'status': w['research_status']} for w in words]
    atomic_json(DATA / 'queue.json', queue)
    (ROOT / 'PROJECT_STATE.md').write_text(
        '# Состояние проекта\n\n'
        f'Корпус SHA-256: `{manifest["corpus_sha256"]}`. Профиль: `{PROFILE}`.\n\n'
        f'Аятов: {manifest["numbered_verses"]}; токенов: {manifest["token_count"]}; '
        f'уникальных словоформ: {manifest["unique_words"]}.\n\n'
        'Исследование по статусам:\n\n' + ''.join(f'- {k}: {v}\n' for k, v in sorted(status.items())) +
        '\nАвторитетное состояние: отдельные JSON в `data/words/`. Очередь и таблицы пересобираются из них.\n', encoding='utf-8')
    return status


def verify():
    manifest, words = read_words()
    _, verses, _, _ = load_corpus()
    source_verses = {(s, a): text.split() for s, a, text in verses}
    indexed = Counter()
    positions = set()
    verse_counts = Counter()
    verse_positions = defaultdict(set)
    with (DATA / 'occurrences.csv').open(encoding='utf-8', newline='') as handle:
        for row in csv.DictReader(handle):
            key = (int(row['sura']), int(row['ayah']), int(row['position']))
            if key in positions:
                raise ValueError(f'Duplicate occurrence {key}')
            positions.add(key)
            verse_counts[key[:2]] += 1
            verse_positions[key[:2]].add(key[2])
            source_pos = int(row['source_position'])
            if source_verses[key[:2]][source_pos - 1] != row['source_token']:
                raise ValueError(f'Source offset mismatch {key}')
            if normalize(row['source_token']) != (row['vocalized_nfc'], row['word']):
                raise ValueError(f'Normalization mismatch {key}')
            indexed[row['word_id']] += 1
    assert len(positions) == manifest['token_count']
    assert len(verse_counts) == manifest['numbered_verses']
    assert all(verse_positions[verse] == set(range(1, count + 1)) for verse, count in verse_counts.items())
    assert all(indexed[w['word_id']] == w['frequency'] for w in words)
    assert sum(w['frequency'] for w in words) == manifest['token_count']
    assert digest((DATA / 'occurrences.csv').read_bytes()) == manifest['occurrences_sha256']
    with (DATA / 'unnumbered_initial_basmalas.csv').open(encoding='utf-8', newline='') as handle:
        extra = list(csv.DictReader(handle))
    assert len(extra) == manifest['unnumbered_initial_basmala_tokens_excluded'] == 448
    assert len({(int(r['sura']), int(r['source_position'])) for r in extra}) == 448
    assert all(source_verses[(int(r['sura']), 1)][int(r['source_position']) - 1] == r['source_token'] for r in extra)
    return {'words': len(words), 'tokens': len(positions), 'verses': len(verse_counts)}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('command', choices=['build', 'verify', 'export', 'status'])
    args = parser.parse_args()
    result = {'build': build, 'verify': verify, 'export': export, 'status': lambda: Counter(w['research_status'] for w in read_words()[1])}[args.command]()
    print(json.dumps(result, ensure_ascii=False, indent=2))
